"""One-off script used to build data/transactions.xlsx. Not needed to run the app."""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

wb = Workbook()

# ---------------------------------------------------------------------------
# Transactions sheet
# ---------------------------------------------------------------------------
ws = wb.active
ws.title = "Transactions"

headers = ["Date", "Description", "Category", "Amount", "Account"]
header_font = Font(name="Arial", bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="2E5266", end_color="2E5266", fill_type="solid")

for col_idx, header in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")

# One example row so the expected format is obvious. Highlighted so it reads as a sample.
sample_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
sample_row = [date(2026, 1, 15), "Example - Tesco Superstore", "Groceries", 42.15, "Current Account"]
for col_idx, value in enumerate(sample_row, start=1):
    cell = ws.cell(row=2, column=col_idx, value=value)
    cell.fill = sample_fill
    cell.font = Font(name="Arial", italic=True)

ws["A2"].number_format = "yyyy-mm-dd"
ws["D2"].number_format = "#,##0.00"

col_widths = {"A": 14, "B": 34, "C": 18, "D": 12, "E": 20}
for col, width in col_widths.items():
    ws.column_dimensions[col].width = width

ws.freeze_panes = "A2"

# ---------------------------------------------------------------------------
# Budget sheet
# ---------------------------------------------------------------------------
budget_ws = wb.create_sheet("Budget")

budget_headers = ["Category", "Monthly Budget", "Notes"]
for col_idx, header in enumerate(budget_headers, start=1):
    cell = budget_ws.cell(row=1, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")

# Starter budget matching the example categories above. Edit freely - add,
# remove, or rename rows to match the categories you actually use.
starter_budget = [
    ("Groceries", 300, ""),
    ("Eating Out", 120, ""),
    ("Transport", 100, ""),
    ("Subscriptions", 40, ""),
    ("Utilities", 150, ""),
    ("Rent", 950, ""),
    ("Shopping", 100, ""),
    ("Entertainment", 60, "Example row - edit or delete freely"),
]
for row_idx, (cat, amount, note) in enumerate(starter_budget, start=2):
    budget_ws.cell(row=row_idx, column=1, value=cat)
    c = budget_ws.cell(row=row_idx, column=2, value=amount)
    c.number_format = "#,##0.00"
    budget_ws.cell(row=row_idx, column=3, value=note)

budget_ws.column_dimensions["A"].width = 18
budget_ws.column_dimensions["B"].width = 16
budget_ws.column_dimensions["C"].width = 40
budget_ws.freeze_panes = "A2"

# ---------------------------------------------------------------------------
# Read Me sheet (legend)
# ---------------------------------------------------------------------------
notes = wb.create_sheet("Read Me")
notes.column_dimensions["A"].width = 100

lines = [
    ("How to use this file", True),
    ("", False),
    ("1. Add one row per transaction on the 'Transactions' tab, directly under the existing rows.", False),
    ("2. Keep the highlighted example row (row 2) as a formatting reference, or delete it once you have your own data - it's not required.", False),
    ("3. Columns:", False),
    ("   - Date: the transaction date (any recognisable date format).", False),
    ("   - Description: merchant / payee / memo text.", False),
    ("   - Category: e.g. Groceries, Rent, Transport, Eating Out, Subscriptions, Utilities, Travel, Salary.", False),
    ("   - Amount: enter spending as a POSITIVE number and income/refunds as a NEGATIVE number.", False),
    ("     (If your bank statements do the opposite, just tick the 'negative = spending' box in the dashboard sidebar.)", False),
    ("   - Account: which bank account or card the transaction is on.", False),
    ("4. Save the file, then click 'Reload data' in the dashboard sidebar (or just refresh the browser tab).", False),
    ("", False),
    ("Setting a budget", True),
    ("", False),
    ("5. On the 'Budget' tab, enter one row per category with the amount you want to spend per month.", False),
    ("   A few starter rows are included - edit the amounts, and add or delete rows to match your own categories.", False),
    ("6. Use the 'Notes' column for anything worth remembering, e.g. 'going up after rent renewal in March'.", False),
    ("7. Open the dashboard's 'Budget vs Actual' page (in the page menu) to see how you're tracking each month,", False),
    ("   which categories are over, and exactly which transactions are driving it.", False),
    ("8. You can also edit the budget from inside that page and click 'Save budget to Excel' - it writes back", False),
    ("   to this same file, so both routes always stay in sync.", False),
    ("", False),
    ("Do not rename the 'Transactions' or 'Budget' sheets or their column headers - the dashboard reads them by name.", False),
]

for row_idx, (text, bold) in enumerate(lines, start=1):
    cell = notes.cell(row=row_idx, column=1, value=text)
    cell.font = Font(name="Arial", bold=bold, size=13 if bold else 11)

wb.save("data/transactions.xlsx")
print("Wrote data/transactions.xlsx")
